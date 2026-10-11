from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "portfolio-summary"
OUTPUT = OUT_DIR / "白龙海-自荐信.docx"
FONT_BODY = "SimSun"
FONT_HEADING = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(95, 95, 95)


def set_font(run, name, size, bold=False, color=BLACK):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), name)
    return run


def set_spacing(paragraph, before=0, after=0, line=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, content):
    paragraph = doc.add_paragraph(style="Letter Body")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    set_font(paragraph.add_run(content), FONT_BODY, 11)
    return paragraph


def add_letter():
    doc = Document()
    section = doc.sections[0]
    # The supplied screenshot is the visual authority: A4, centered title,
    # restrained black typography, spacious formal-letter rhythm.
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.5

    if "Letter Body" not in doc.styles:
        doc.styles.add_style("Letter Body", WD_STYLE_TYPE.PARAGRAPH)
    body_style = doc.styles["Letter Body"]
    body_style.base_style = normal
    body_style.font.name = FONT_BODY
    body_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    body_style.font.size = Pt(11)
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(8)
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.widow_control = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(title, before=0, after=18, line=1.0)
    set_font(title.add_run("自荐信"), FONT_HEADING, 22, bold=True)

    salutation = doc.add_paragraph()
    set_spacing(salutation, after=10, line=1.4)
    set_font(salutation.add_run("尊敬的招聘负责人："), FONT_BODY, 11, bold=True)

    add_body(doc, "您好！感谢您在百忙之中阅读我的自荐信。我叫白龙海，目前就读于电子科技大学成都学院动画专业本科，希望应聘贵公司在上海地区的 3D 设计师或相关岗位。")

    add_body(doc, "我拥有三年以上三维创作经验，目前已整理完成 58 项作品。我的创作方向以硬表面建模、机械结构设计和三维视觉表现为主，能够使用 Blender、Plasticity 等软件完成从概念分析、形体搭建、拓扑与布线、展 UV、PBR 材质、灯光渲染到后期输出的完整流程。")

    add_body(doc, "在项目实践中，我完成了战术机甲、人形机甲、枪械、科幻摩托等多类作品，也进行了机械结构绑定与控制器设置。制作过程中，我不仅关注模型外形是否准确，也重视部件之间的结构关系、功能逻辑和视觉节奏；面对复杂曲面与细节时，我会通过拆解参考、规划层级和反复验证来保持模型的可信度与可编辑性。")

    add_body(doc, "除静态模型外，我也具备独立动画与影像制作能力，能够完成前期策划、分镜设计、骨骼绑定、动作设计、镜头调度、引擎渲染、后期合成与剪辑。在校期间，我曾担任校园广播电视台台长，带领团队完成多项融媒体中心拍摄任务，积累了项目统筹、沟通协作和现场执行经验。参与制作的《错位时空》获得人民日报转发，并获得第六届四川省高校大学生微电影二等奖。")

    add_body(doc, "我熟悉 Blender、Plasticity、Substance 3D Painter、Marvelous Designer、Spine，以及 Adobe Photoshop、Adobe Premiere Pro、Adobe After Effects、DaVinci Resolve 等软件。面对陌生工具时，我能够快速理解底层逻辑并上手操作；遇到制作问题时，我会先分析原因、拆分任务并逐项验证，保持对细节和完成质量的要求。")

    add_body(doc, "我希望进入更专业的项目环境，将个人积累转化为稳定、可靠的制作成果，并在真实项目中继续提升建模、动画与视觉表达能力。若有机会，我可在一个月内到岗。随信附上个人简历与作品集，期待进一步介绍我的项目过程和创作思路。感谢您的时间与考虑，诚盼回复！")

    closing = doc.add_paragraph()
    set_spacing(closing, before=4, after=0, line=1.5)
    set_font(closing.add_run("此致"), FONT_BODY, 11)
    respect = doc.add_paragraph()
    respect.paragraph_format.left_indent = Cm(0.74)
    set_spacing(respect, after=16, line=1.5)
    set_font(respect.add_run("敬礼！"), FONT_BODY, 11, bold=True)

    for value, bold in [
        ("自荐人：白龙海", True),
        ("邮箱：2052130377@qq.com", False),
        ("微信：blh2052130377", False),
        ("2026 年 9 月 1 日", False),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_spacing(paragraph, after=1, line=1.15)
        set_font(paragraph.add_run(value), FONT_BODY, 10.5, bold=bold, color=BLACK if bold else GRAY)

    doc.core_properties.title = "白龙海 - 自荐信"
    doc.core_properties.subject = "3D 建模师 / 三维设计相关岗位自荐信"
    doc.core_properties.author = "白龙海"
    doc.core_properties.last_modified_by = "白龙海"
    doc.core_properties.keywords = "3D建模, 自荐信, 硬表面建模, Blender, Plasticity"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


def audit(output):
    doc = Document(output)
    assert len(doc.sections) == 1 and not doc.tables
    section = doc.sections[0]
    assert abs(section.page_width - Cm(21)) < 1000 and abs(section.page_height - Cm(29.7)) < 1000
    assert abs(section.top_margin - Cm(1.8)) < 1000 and abs(section.bottom_margin - Cm(1.8)) < 1000
    assert abs(section.left_margin - Cm(2.4)) < 1000 and abs(section.right_margin - Cm(2.4)) < 1000
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "白龙海", "电子科技大学成都学院", "3D 设计师", "58 项作品",
        "战术机甲", "人形机甲", "枪械", "科幻摩托", "机械结构绑定",
        "动画专业本科", "校园广播电视台台长", "微电影二等奖", "人民日报转发", "一个月内到岗",
        "Blender", "Plasticity", "Substance 3D Painter", "DaVinci Resolve",
        "2052130377@qq.com", "blh2052130377",
    ]
    for value in required:
        assert value in text, value
    for value in ["XXX", "待补充", "TODO", "TBD", "全民简历", "工商管理", "行政人事", "上海养擎"]:
        assert value not in text, value
    assert len(text.replace("\n", "")) >= 750
    with ZipFile(output) as package:
        xml = package.read("word/document.xml").decode("utf-8")
        assert "�" not in xml and "<w:tbl>" not in xml
        assert "w:type=\"page\"" not in xml
        assert "SimSun" in xml and "Microsoft YaHei" in xml
    return {"paragraphs": len(doc.paragraphs), "characters": len(text.replace("\n", "")), "tables": len(doc.tables), "placeholders": False}


if __name__ == "__main__":
    output = add_letter()
    print(output)
    print(audit(output))
