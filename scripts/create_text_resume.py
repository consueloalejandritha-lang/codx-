from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "portfolio-summary"
OUTPUT = OUTPUT_DIR / "白龙海-文字版简历.docx"

FONT = "Microsoft YaHei"
BLACK = RGBColor(5, 5, 5)
MUTED = RGBColor(92, 92, 92)
LIGHT = "D9D9D9"
CONTENT_WIDTH = Inches(7.18)

# Design system: compact_reference_guide + customer_pack-style left header.
# Named resume_one_page overrides: 0.55/0.66 in margins, 9.6 pt body,
# 1.10 body leading, 11.5 pt section headings, and no running header.


def set_font(run, size=None, bold=None, color=BLACK, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color
    return run


def set_spacing(paragraph, before=0, after=0, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_bottom_border(paragraph, color=LIGHT, size=6, space=3):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_section_heading(doc, title):
    paragraph = doc.add_paragraph(style="Resume Section")
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(title), size=11.5, bold=True)
    add_bottom_border(paragraph)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="Resume Bullet")
    set_font(paragraph.add_run(text), size=9.6)
    return paragraph


def add_project(doc, title, english, year, description):
    heading = doc.add_paragraph(style="Project Heading")
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    set_font(heading.add_run(title), size=10.5, bold=True)
    set_font(heading.add_run(f"  {english}"), size=8.2, color=MUTED)
    if year:
        set_font(heading.add_run(f"\t{year}"), size=8.8, bold=True, color=MUTED)
    add_bullet(doc, description)


def add_real_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(element.get(qn("w:abstractNumId"))) for element in numbering.findall(qn("w:abstractNum"))]
    existing_nums = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abstract, default=0) + 1
    num_id = max(existing_nums, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); level.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "bullet"); level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "•"); level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc"); lvl_jc.set(qn("w:val"), "left"); level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "432"); tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "432"); ind.set(qn("w:hanging"), "230"); p_pr.append(ind)
    spacing = OxmlElement("w:spacing"); spacing.set(qn("w:before"), "0"); spacing.set(qn("w:after"), "20"); spacing.set(qn("w:line"), "264"); spacing.set(qn("w:lineRule"), "auto"); p_pr.append(spacing)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), FONT); fonts.set(qn("w:hAnsi"), FONT); fonts.set(qn("w:eastAsia"), FONT); r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId"); abstract_ref.set(qn("w:val"), str(abstract_id)); num.append(abstract_ref)
    numbering.append(num)
    style = doc.styles["Resume Bullet"]
    p_pr = style.element.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId"); num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    p_pr.append(num_pr)


def build_resume():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True

    for name in ["Resume Section", "Resume Bullet", "Project Heading"]:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.base_style = normal
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.styles["Resume Section"].paragraph_format.space_before = Pt(7)
    doc.styles["Resume Section"].paragraph_format.space_after = Pt(4)
    doc.styles["Resume Section"].paragraph_format.line_spacing = 1.0
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    doc.styles["Resume Section"].element.get_or_add_pPr().append(outline)
    doc.styles["Resume Bullet"].paragraph_format.space_after = Pt(1)
    doc.styles["Resume Bullet"].paragraph_format.line_spacing = 1.1
    doc.styles["Project Heading"].paragraph_format.space_before = Pt(3)
    doc.styles["Project Heading"].paragraph_format.space_after = Pt(1)
    doc.styles["Project Heading"].paragraph_format.line_spacing = 1.0
    add_real_bullet_numbering(doc)

    name = doc.add_paragraph()
    set_spacing(name, after=0, line=1.0)
    set_font(name.add_run("白龙海"), size=27, bold=True)

    role = doc.add_paragraph()
    set_spacing(role, after=3, line=1.0)
    set_font(role.add_run("3D 建模师  |  独立动画创作者"), size=12.5, bold=True, color=MUTED)

    contact = doc.add_paragraph()
    set_spacing(contact, after=5, line=1.0)
    set_font(contact.add_run("邮箱：2052130377@qq.com   |   微信：blh2052130377   |   电子科技大学成都学院"), size=9.2, color=MUTED)
    add_bottom_border(contact, color="050505", size=10, space=5)

    add_section_heading(doc, "个人简介  PROFILE")
    summary = doc.add_paragraph()
    set_spacing(summary, after=1, line=1.16)
    set_font(summary.add_run(
        "专注于硬表面建模、机械结构设计与三维视觉创作，具备从概念分析、模型制作、拓扑与 UV、材质灯光到动画呈现的完整制作能力。创作中重视结构关系、功能逻辑与视觉节奏，追求准确的造型、克制的细节和具有辨识度的动态表达。"
    ), size=9.6)

    add_section_heading(doc, "核心能力  CORE SKILLS")
    for item in [
        "硬表面建模：根据功能与结构关系拆分形体，在机械逻辑、造型比例和细节密度之间取得平衡。",
        "拓扑与布线：面向细分、变形与动画需求规划拓扑，使复杂曲面保持干净、稳定并便于调整。",
        "全流程制作：能够独立完成建模、拓扑、展 UV、PBR 材质、灯光、渲染及后期合成。",
        "独立动画制作：完成分镜、机械绑定、动作设计、镜头调度、渲染与剪辑，形成连贯的动态表达。",
    ]:
        add_bullet(doc, item)

    add_section_heading(doc, "代表项目  SELECTED PROJECTS")
    add_project(doc, "战术机甲", "TACTICAL MECH", "2026", "围绕战术装备与机械功能完成硬表面造型、结构细化和最终视觉呈现，重点处理大型结构之间的比例、层级与细节密度。")
    add_project(doc, "人形机甲", "MECH FORM STUDY", "2026", "以人形机甲为主题进行复杂机械结构与形体语言研究，完成主要结构拆分、硬表面细化及模型展示。")
    add_project(doc, "枪械制造", "FIREARM MODELING", "2025", "使用 Plasticity 的 NURBS / CAD 工作流构建锐利、干净的工业曲面，并完成材质、灯光与成品展示。")
    add_project(doc, "机械结构绑定与控制", "RIGGING & CONTROLLER SETUP", "", "针对机械结构建立绑定与控制器方案，梳理部件层级和运动关系，并通过动画与动态演示验证控制逻辑。")

    add_section_heading(doc, "软件技能  SOFTWARE")
    software = [
        ("三维建模与材质", "Blender、Plasticity、Substance 3D Painter、Marvelous Designer、DAZ Studio、酷家乐"),
        ("动画与绑定", "Spine、Adobe After Effects"),
        ("后期与设计", "Adobe Photoshop、Adobe Premiere Pro、Adobe Media Encoder、DaVinci Resolve、剪映"),
    ]
    for label, value in software:
        p = doc.add_paragraph()
        set_spacing(p, after=1, line=1.08)
        set_font(p.add_run(f"{label}："), size=9.4, bold=True)
        set_font(p.add_run(value), size=9.4)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_spacing(footer_p, after=0, line=1.0)
    set_font(footer_p.add_run("BAILONGHAI  /  3D ARTIST"), size=7.5, color=RGBColor(145, 145, 145))

    doc.core_properties.title = "白龙海 - 文字版简历"
    doc.core_properties.subject = "3D 建模师 / 独立动画创作者"
    doc.core_properties.author = "白龙海"
    doc.core_properties.last_modified_by = "白龙海"
    doc.core_properties.keywords = "3D建模, 硬表面建模, Plasticity, Blender, 动画"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    result = build_resume()
    print(result)
