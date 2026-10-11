from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
RESUME = ROOT / "portfolio-summary" / "白龙海-文字版简历.docx"
assert RESUME.exists() and RESUME.stat().st_size > 20_000
doc = Document(RESUME)
section = doc.sections[0]
assert section.page_width == 7_772_400 and section.page_height == 10_058_400  # Letter
assert section.top_margin == 502_920 and section.bottom_margin == 502_920
assert section.left_margin == 603_250 and section.right_margin == 603_250
assert len(doc.sections) == 1 and not doc.tables
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
for required in [
    "白龙海", "3D 建模师", "个人简介", "核心能力", "代表项目", "软件技能",
    "战术机甲", "人形机甲", "枪械制造", "机械结构绑定与控制",
    "Blender", "Plasticity", "Substance 3D Painter", "DaVinci Resolve",
    "2052130377@qq.com", "blh2052130377", "电子科技大学成都学院",
]:
    assert required in text, required
for forbidden in ["待补充", "TODO", "TBD", "Lorem", "Codex", "OpenAI"]:
    assert forbidden not in text
assert all(not paragraph.text.lstrip().startswith(("•", "- ")) for paragraph in doc.paragraphs)
for index in range(6, 10):
    num_pr = doc.paragraphs[index].style.element.pPr.numPr
    assert num_pr is not None and num_pr.numId is not None
for index in [12, 14, 16, 18]:
    assert doc.paragraphs[index].style.element.pPr.numPr is not None
assert doc.styles["Resume Section"].element.pPr.find(qn("w:outlineLvl")) is not None
with ZipFile(RESUME) as package:
    names = set(package.namelist())
    assert {"word/document.xml", "word/styles.xml", "word/numbering.xml"} <= names
    document_xml = package.read("word/document.xml").decode("utf-8")
    styles_xml = package.read("word/styles.xml").decode("utf-8")
    numbering_xml = package.read("word/numbering.xml").decode("utf-8")
    assert "�" not in document_xml
    assert "<w:tbl>" not in document_xml
    assert "w:type=\"page\"" not in document_xml
    assert "Microsoft YaHei" in styles_xml
    assert "w:lvlText w:val=\"•\"" in numbering_xml
    assert "w:hanging=\"230\"" in numbering_xml and "w:left=\"432\"" in numbering_xml
print({"file": str(RESUME), "paragraphs": len(doc.paragraphs), "sections": len(doc.sections), "tables": len(doc.tables), "placeholders": False, "real_bullets": True})
